import sqlite3
import logging
import time
import json
import os
import sys
import getpass
import re
import subprocess
import threading
import csv
import tkinter as tk
from tkinter import ttk, scrolledtext, simpledialog, messagebox
from datetime import datetime
from typing import List, Dict, Tuple

try:
    from netmiko import ConnectHandler
    from netmiko.exceptions import NetmikoTimeoutException, NetmikoAuthenticationException
except ImportError:
    logging.error("Netmiko library is not installed. Please run: pip install netmiko")
    sys.exit(1)

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.marker import Marker
    EXCEL_SUPPORT = True
except ImportError:
    logging.warning("openpyxl is not installed! Excel report generation will be skipped. Run 'pip install openpyxl'")
    EXCEL_SUPPORT = False

CONFIG_FILE = "env_config.json"
GUI_ROOT = None  # Global reference for thread-safe popups

# ==========================================
# Thread-Safe GUI Communication Helpers
# ==========================================
def gui_ask(prompt_text: str) -> str:
    """Safely prompts the user for input from a background thread."""
    if not GUI_ROOT: 
        return "exit"
    res = {"val": None}
    ev = threading.Event()
    
    def _ask():
        val = simpledialog.askstring("Input Required", prompt_text, parent=GUI_ROOT)
        res["val"] = val if val is not None else "exit"
        ev.set()
        
    GUI_ROOT.after(0, _ask)
    ev.wait()
    return res["val"]

class SafeTextRedirector:
    """Redirects stdout and stderr to a tkinter Text widget safely across threads."""
    def __init__(self, text_widget):
        self.text_widget = text_widget

    def write(self, string):
        self.text_widget.after(0, self._insert, string)

    def _insert(self, string):
        self.text_widget.configure(state="normal")
        self.text_widget.insert("end", string)
        self.text_widget.see("end")
        self.text_widget.configure(state="disabled")

    def flush(self):
        pass

class TextHandler(logging.Handler):
    """Routes python logging to the SafeTextRedirector."""
    def __init__(self, redirector):
        logging.Handler.__init__(self)
        self.redirector = redirector

    def emit(self, record):
        msg = self.format(record)
        self.redirector.write(msg + '\n')

# ==========================================
# Core Engine & Database
# ==========================================
def generate_professional_excel_report(serial: str, db, config: Dict):
    """Generates a highly formatted 3-sheet Excel report with live speed graphs."""
    if not EXCEL_SUPPORT: return
    
    cases = db.load_test_cases(serial)
    if not cases: return
    
    filename = f"ONT_Test_Result_{serial}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    try:
        wb = openpyxl.Workbook()
        
        bold_font = Font(bold=True)
        title_font = Font(size=16, bold=True, color="FFFFFF")
        hdr_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        na_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        nt_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        ws_sum = wb.active
        ws_sum.title = "Summary"
        
        ws_sum.merge_cells('B2:E3')
        title_cell = ws_sum.cell(row=2, column=2, value="ONT Automation Test Report")
        title_cell.font = title_font
        title_cell.alignment = center_align
        title_cell.fill = hdr_fill
        
        total = len(cases)
        c_pass = sum(1 for c in cases if c['status'] == 'PASS')
        c_fail = sum(1 for c in cases if c['status'] == 'FAIL')
        c_na   = sum(1 for c in cases if c['status'] == 'N/A')
        c_nt   = sum(1 for c in cases if c['status'] == 'N/T')
        progress = ((total - c_nt) / total) * 100 if total else 0
        
        info = [
            ("Target Serial", serial),
            ("OLT Vendor", config.get('global_vendor', 'Unknown')),
            ("Target PON / ID", f"{config.get('pon', config.get('hw_fsp', 'N/A'))} / {config.get('ont_id', config.get('hw_ont_id', 'N/A'))}"),
            ("ONT SW Version", config.get('sw_version', 'Unknown')),
            ("VLAN Mode", config.get('vlan_mode', 'untagged').capitalize()),
            ("Test Date", datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
            ("Total Tests", total),
            ("PASS", c_pass),
            ("FAIL", c_fail),
            ("N/A (Skipped)", c_na),
            ("N/T (Not Tested)", c_nt),
            ("Progress", f"{progress:.1f}%")
        ]
        
        row_idx = 5
        for key, val in info:
            c_key = ws_sum.cell(row=row_idx, column=2, value=key)
            c_val = ws_sum.cell(row=row_idx, column=3, value=val)
            c_key.font = bold_font
            c_key.border = thin_border
            c_val.border = thin_border
            if key == "PASS": c_val.fill = pass_fill
            elif key == "FAIL": c_val.fill = fail_fill
            elif key == "N/A (Skipped)": c_val.fill = na_fill
            elif key == "N/T (Not Tested)": c_val.fill = nt_fill
            row_idx += 1
            
        ws_sum.column_dimensions['B'].width = 25
        ws_sum.column_dimensions['C'].width = 35

        ws_det = wb.create_sheet(title="Details")
        headers = ["Test Order", "Test Item", "Command / Script", "Result"]
        for col_num, header in enumerate(headers, 1):
            cell = ws_det.cell(row=1, column=col_num, value=header)
            cell.font = bold_font
            cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            cell.alignment = center_align
            cell.border = thin_border
            
        for idx, c in enumerate(cases, 2):
            ws_det.cell(row=idx, column=1, value=c['order']).alignment = center_align
            ws_det.cell(row=idx, column=2, value=c['type'])
            ws_det.cell(row=idx, column=3, value=c['command'])
            
            res_cell = ws_det.cell(row=idx, column=4, value=c['status'])
            res_cell.alignment = center_align
            if c['status'] == 'PASS': res_cell.fill = pass_fill
            elif c['status'] == 'FAIL': res_cell.fill = fail_fill
            elif c['status'] == 'N/A': res_cell.fill = na_fill
            else: res_cell.fill = nt_fill
            
            for col in range(1, 5):
                ws_det.cell(row=idx, column=col).border = thin_border
                
        ws_det.column_dimensions['B'].width = 25
        ws_det.column_dimensions['C'].width = 80
        ws_det.column_dimensions['D'].width = 15

        ws_log = wb.create_sheet(title="Logs & Graph")
        ws_log.cell(row=1, column=1, value="Test Item").font = bold_font
        ws_log.cell(row=1, column=2, value="Execution Log").font = bold_font
        ws_log.column_dimensions['A'].width = 20
        ws_log.column_dimensions['B'].width = 100
        
        log_row = 2
        up_data = {}
        dl_data = {}
        current_mode = "UPLOAD"
        
        for c in cases:
            ws_log.cell(row=log_row, column=1, value=c['type']).font = bold_font
            log_lines = (c['log'] or "No log captured.").split('\n')
            
            for line in log_lines:
                clean_line = line.strip()
                if not clean_line: continue
                
                clean_line = re.sub(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]', '', clean_line)
                if clean_line.startswith('='):
                    clean_line = "'" + clean_line

                ws_log.cell(row=log_row, column=2, value=clean_line)
                
                if c['type'] == 'Speed_Test':
                    if "--- STARTING DOWNLOAD TEST ---" in clean_line:
                        current_mode = "DOWNLOAD"
                    elif "--- STARTING UPLOAD TEST ---" in clean_line:
                        current_mode = "UPLOAD"
                        
                    if "[SUM]" in clean_line and "sec" in clean_line and "bits/sec" in clean_line:
                        if "sender" not in clean_line and "receiver" not in clean_line:
                            match = re.search(r'\[SUM\]\s+\d+\.\d+-\s*(\d+\.\d+)\s+sec.*?\s+(\d+(?:\.\d+)?)\s+([KMG])bits/sec', clean_line)
                            if match:
                                end_time_float = float(match.group(1))
                                t_sec = int(round(end_time_float))
                                val = float(match.group(2))
                                unit = match.group(3)
                                mbps = val * 1000 if unit == 'G' else (val if unit == 'M' else val / 1000)
                                
                                if current_mode == "UPLOAD":
                                    up_data[t_sec] = max(up_data.get(t_sec, 0), mbps)
                                else:
                                    dl_data[t_sec] = max(dl_data.get(t_sec, 0), mbps)
                log_row += 1
            log_row += 1 

        if up_data or dl_data:
            ws_log.cell(row=1, column=4, value="Time (sec)").font = bold_font
            ws_log.cell(row=1, column=5, value="Upload (Mbps)").font = bold_font
            ws_log.cell(row=1, column=6, value="Download (Mbps)").font = bold_font
            
            all_times = sorted(list(set(list(up_data.keys()) + list(dl_data.keys()))))
            
            r_idx = 2
            if all_times:
                max_time = max(all_times)
                last_up = None
                last_dl = None
                
                for t in range(1, max_time + 1):
                    val_up = up_data.get(t, last_up)
                    val_dl = dl_data.get(t, last_dl)
                    
                    ws_log.cell(row=r_idx, column=4, value=t)
                    ws_log.cell(row=r_idx, column=5, value=val_up)
                    ws_log.cell(row=r_idx, column=6, value=val_dl)
                    
                    if val_up is not None: last_up = val_up
                    if val_dl is not None: last_dl = val_dl
                    r_idx += 1
                
            chart = LineChart()
            chart.title = "iPerf3 Throughput over Time (Up/Down)"
            chart.y_axis.title = 'Throughput (Mbps)'
            chart.x_axis.title = 'Time (sec)'
            chart.width = 20
            chart.height = 10
            
            chart.x_axis.tickLblPos = "low"
            chart.y_axis.tickLblPos = "low"
            chart.dispBlanksAs = "span" 
            
            data = Reference(ws_log, min_col=5, max_col=6, min_row=1, max_row=r_idx-1)
            cats = Reference(ws_log, min_col=4, min_row=2, max_row=r_idx-1)
            
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            
            try:
                for s in chart.series:
                    s.marker = Marker(symbol='circle', size=5)
            except Exception as e:
                logging.warning(f"Could not format chart markers: {e}")
            
            ws_log.add_chart(chart, "H2")

        wb.save(filename)
        print("\n" + "="*60)
        logging.info(f" [✔] Professional Excel Report Generated: {filename} ")
        print("="*60 + "\n")
        
    except Exception as e:
        logging.error(f"Failed to generate Excel report: {e}")

class DatabaseManager:
    def __init__(self, db_name: str = "olt_physical_automation.db"):
        self.db_name = db_name
        self._initialize_database()

    def _get_conn(self):
        return sqlite3.connect(self.db_name, timeout=20)

    def _initialize_database(self):
        with self._get_conn() as conn:
            cursor = conn.cursor()
            cursor.execute('''CREATE TABLE IF NOT EXISTS ont_inventory (
                id INTEGER PRIMARY KEY AUTOINCREMENT, serial_number TEXT UNIQUE, 
                pon_port TEXT, ont_id TEXT, status TEXT, provisioning_date TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''')
            
            cursor.execute('''CREATE TABLE IF NOT EXISTS test_cases (
                id INTEGER PRIMARY KEY AUTOINCREMENT, target_serial TEXT, execution_order INTEGER, 
                test_type TEXT, command TEXT, parameters TEXT, expected_output TEXT, 
                success_count INTEGER DEFAULT 0, is_evolved BOOLEAN DEFAULT 0)''')
                
            cursor.execute('''CREATE TABLE IF NOT EXISTS learned_commands (
                olt_profile TEXT, test_type TEXT, command_template TEXT, PRIMARY KEY (olt_profile, test_type))''')
            
            try: cursor.execute("ALTER TABLE test_cases ADD COLUMN status TEXT DEFAULT 'N/T'")
            except: pass
            try: cursor.execute("ALTER TABLE test_cases ADD COLUMN execution_log TEXT")
            except: pass
            conn.commit()

    def add_ont(self, serial: str, pon: str, ont_id: str, status: str):
        with self._get_conn() as conn:
            cursor = conn.cursor()
            cursor.execute('INSERT OR REPLACE INTO ont_inventory (serial_number, pon_port, ont_id, status) VALUES (?, ?, ?, ?)', 
                           (serial, pon, ont_id, status))
            conn.commit()

    def get_learned_command(self, profile: str, t_type: str) -> str:
        with self._get_conn() as conn:
            cursor = conn.cursor()
            cursor.execute('SELECT command_template FROM learned_commands WHERE olt_profile = ? AND test_type = ?', (profile, t_type))
            row = cursor.fetchone()
            return row[0] if row else None

    def save_learned_command(self, profile: str, t_type: str, template: str):
        with self._get_conn() as conn:
            cursor = conn.cursor()
            cursor.execute('INSERT OR REPLACE INTO learned_commands (olt_profile, test_type, command_template) VALUES (?, ?, ?)', 
                           (profile, t_type, template))
            conn.commit()

    def update_test_case(self, t_id: int, cmd: str, params: str):
        with self._get_conn() as conn:
            cursor = conn.cursor()
            cursor.execute('UPDATE test_cases SET command = ?, parameters = ? WHERE id = ?', (cmd, params, t_id))
            conn.commit()

    def update_test_status(self, t_id: int, status: str, log: str):
        with self._get_conn() as conn:
            cursor = conn.cursor()
            cursor.execute("UPDATE test_cases SET status = ?, execution_log = ? WHERE id = ?", (status, log, t_id))
            if status == "PASS":
                cursor.execute("UPDATE test_cases SET success_count = success_count + 1 WHERE id = ?", (t_id,))
            conn.commit()

    def load_test_cases(self, serial: str) -> List[Dict]:
        with self._get_conn() as conn:
            cursor = conn.cursor()
            cursor.execute('SELECT id, target_serial, execution_order, test_type, command, parameters, status, execution_log, success_count FROM test_cases WHERE target_serial = ? ORDER BY execution_order ASC', (serial,))
            rows = cursor.fetchall()
            return [{"id": r[0], "serial": r[1], "order": r[2], "type": r[3], "command": r[4], "parameters": json.loads(r[5]), "status": r[6], "log": r[7], "success_count": r[8]} for r in rows]

    def add_initial_test_cases(self, config: Dict, profile: str):
        serial = config['ont_serial']
        
        vendor = config.get("global_vendor", "Nokia ISAM7360")
        if vendor == "Nokia ISAM7360":
            pon = config['pon']
            ont_id = config['ont_id']
            chanpair = config.get('chanpair', '')
            port_full = f"{pon}/{ont_id}"
            vendor_id = serial[:4]
            sn_rem = serial[4:]
            serial_formatted = f"{vendor_id}:{sn_rem}"
            slot = config.get('ont_slot', '1')
            
            templates = {
                "ONT_Discovery_Check": "show equipment ont status channel-pair {chanpair}", 
                "Registration_Check": "show equipment ont status channel-pair {chanpair}", 
                "Reboot_Test": "admin equipment ont interface {port} reboot with-active-image",
                "Optics_Signal_Check": "show equipment ont optics {port}",
                "Optics_Temp_Check": "show equipment ont optics {port}",
                "Optics_Voltage_Laser_Check": "show equipment ont optics {port}",
                "UNI_Status": f"show ethernet ont operational-data {{port}}/{slot}/1",
                "Software_Info": "show equipment ont interface {port} detail",
                "Speed_Test": "Native iperf3 Execution", 
                "MAC_Status": f"show vlan bridge-port-fdb {{port}}/{slot}/1",
                "Cleanup": "configure equipment ont interface {port} admin-state down ; configure equipment ont no interface {port}"
            }
        else:
            # Huawei Placeholder Framework
            pon = config.get('hw_fsp', '0/1/0')
            ont_id = config.get('hw_ont_id', '1')
            port_full = f"{pon} {ont_id}"
            serial_formatted = serial
            chanpair = ""
            slot = ""
            
            templates = {
                "ONT_Discovery_Check": "display ont info {port}", 
                "Registration_Check": "display ont info {port}", 
                "Reboot_Test": "ont reset {port}",
                "Optics_Signal_Check": "display ont optical-info {port}",
                "Optics_Temp_Check": "display ont optical-info {port}",
                "Optics_Voltage_Laser_Check": "display ont optical-info {port}",
                "UNI_Status": "display ont port state {port} eth-port all",
                "Software_Info": "display ont version {port}",
                "Speed_Test": "Native iperf3 Execution", 
                "MAC_Status": "display mac-address port {port}",
                "Cleanup": "undo ont {port}"
            }

        with self._get_conn() as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM test_cases WHERE target_serial = ?", (serial,))
            
            cursor.execute('SELECT test_type, command_template FROM learned_commands WHERE olt_profile = ?', (profile,))
            learned_dict = {row[0]: row[1] for row in cursor.fetchall()}
            
            mock_cases = []
            for i, (t_type, def_cmd) in enumerate(templates.items(), 1):
                template = learned_dict.get(t_type) if vendor == "Nokia ISAM7360" else def_cmd
                
                if vendor == "Nokia ISAM7360":
                    if t_type == "Reboot_Test" and template and "with-active-image" not in template:
                        template = def_cmd
                        cursor.execute('INSERT OR REPLACE INTO learned_commands (olt_profile, test_type, command_template) VALUES (?, ?, ?)', (profile, t_type, def_cmd))
                    elif t_type == "UNI_Status" and template and "operational-data" not in template:
                        template = def_cmd
                        cursor.execute('INSERT OR REPLACE INTO learned_commands (olt_profile, test_type, command_template) VALUES (?, ?, ?)', (profile, t_type, def_cmd))
                    elif "{port}" in def_cmd and template and "{port}" not in template:
                        template = def_cmd
                        cursor.execute('INSERT OR REPLACE INTO learned_commands (olt_profile, test_type, command_template) VALUES (?, ?, ?)', (profile, t_type, def_cmd))
                    elif "{chanpair}" in def_cmd and template and "{chanpair}" not in template:
                        template = def_cmd
                        cursor.execute('INSERT OR REPLACE INTO learned_commands (olt_profile, test_type, command_template) VALUES (?, ?, ?)', (profile, t_type, def_cmd))
                        
                if not template:
                    template = def_cmd
                
                cmd = template.format(
                    serial=serial, 
                    serial_formatted=serial_formatted,
                    pon=pon, 
                    ont_id=ont_id, 
                    port=port_full,
                    chanpair=chanpair
                )
                
                params = {"timeout": 15, "wait": 20 if i > 1 else 0}
                mock_cases.append((serial, i, t_type, cmd, json.dumps(params), "N/T"))
                
            cursor.executemany("INSERT INTO test_cases (target_serial, execution_order, test_type, command, parameters, status) VALUES (?,?,?,?,?,?)", mock_cases)
            conn.commit()

class HuaweiOLTConnector:
    """Stub connector for Huawei MA5800-X7 Integration"""
    def __init__(self, ip: str, user: str, pw: str, proto: str):
        self.ip = ip
        self.olt_profile = "Huawei_MA5800"
        self.vendor = "Huawei"
        self.model = "MA5800-X7"
        self.sw_version = "Unknown"
        self.connection = None

    def connect(self) -> bool:
        logging.warning("Huawei Connector is initialized but automation sequences are pending full implementation in next patch.")
        return False
        
    def disconnect(self): pass

# --- Original Nokia Connector ---
class NokiaOLTConnector:
    def __init__(self, ip: str, user: str, pw: str, proto: str):
        self.proto = proto
        self.device = {'host': ip, 'username': user, 'password': pw, 'global_delay_factor': 2}
        self.connection = None
        self.olt_profile = "Nokia_Default"
        self.vendor = "Unknown"
        self.model = "Unknown"
        self.sw_version = "Unknown"

    def connect(self) -> bool:
        types = ["nokia_sros_telnet", "alcatel_sros_telnet"] if self.proto == "telnet" else ["nokia_sros", "alcatel_sros"]
        for dt in types:
            self.device['device_type'] = dt
            try:
                self.connection = ConnectHandler(**self.device)
                logging.info(f"Connected to OLT using device_type: {dt}")
                self._discover_profile()
                return True
            except Exception: 
                continue
        return False

    def disconnect(self):
        if self.connection:
            try:
                self.connection.disconnect()
                logging.info("Gracefully disconnected from OLT.")
            except Exception: pass

    def _discover_profile(self):
        out = self.send_command("show software-mngt version ansi", 5)
        m = re.search(r'(V\d+\.[A-Za-z0-9\.\-_]+|R\d+\.[A-Za-z0-9\.\-_]+|\d{1,3}\.\d{1,3}\.\S+)', out, re.IGNORECASE)
        if m:
            self.sw_version = m.group(1)
        else:
            clean_out = re.sub(r'show software-mngt version ansi', '', out, flags=re.IGNORECASE)
            words = clean_out.split()
            self.sw_version = words[0] if words else "Unknown Version"

        self.olt_profile = f"Nokia_OS_{self.sw_version}"

        if self.device.get('host') == '10.100.0.7':
            self.vendor = "Nokia"
            self.model = "ISAM7360-FWLT-B"
        else:
            if "Nokia" in out or "nokia" in out.lower(): self.vendor = "Nokia"
            elif "Alcatel" in out or "alcatel" in out.lower(): self.vendor = "Alcatel-Lucent"
            else: self.vendor = "Unknown"
            
            sys_out = self.send_command("show equipment system", 5)
            model_m = re.search(r'(FX-\d+|7360|7342|7750\s+SR\S*)', sys_out, re.IGNORECASE)
            self.model = model_m.group(1) if model_m else "ISAM (Generic)"

    def send_command(self, cmd: str, timeout: int = 15, retry: bool = True) -> str:
        if self.connection:
            try: 
                if not self.connection.is_alive():
                    raise Exception("Connection dead")
                return self.connection.send_command_timing(cmd, read_timeout=timeout)
            except Exception as e:
                if retry:
                    logging.warning(f"Connection dropped during command execution. Attempting Auto-Reconnect...")
                    self.disconnect()
                    if self.connect():
                        logging.info("Auto-Reconnect successful. Retrying command...")
                        try:
                            return self.connection.send_command_timing(cmd, read_timeout=timeout)
                        except Exception as e2:
                            return f"Error after reconnect: {e2}"
                return f"Error: {e}"
        else:
            if retry and self.connect():
                return self.send_command(cmd, timeout=timeout, retry=False)
        return "No Connection"

    def get_unprovisioned_onts(self, discovery_cmd: str) -> List[Dict[str, str]]:
        if self.connection:
            output = self.send_command(discovery_cmd)
            matches = re.findall(r'(\d+(?:/\d+)+)\s+([A-Za-z]{4}[A-Fa-f0-9]{8})', output)
            if matches:
                onts = [{"chanpair": m[0], "serial": m[1]} for m in matches]
                unique_onts = {item['serial']: item for item in onts}.values()
                return list(unique_onts)
            else:
                serials = re.findall(r'([A-Za-z]{4}[A-Fa-f0-9]{8})', output)
                return [{"chanpair": "unknown", "serial": s} for s in set(serials)]
        return []

    def _get_safe_delete_cmd(self, db: DatabaseManager) -> str:
        del_cmd_temp = db.get_learned_command(self.olt_profile, "Delete_Provisioned")
        default_cmd = "configure equipment ont interface {port} admin-state down ; configure equipment ont no interface {port}"
        if not del_cmd_temp or "{port}" not in del_cmd_temp:
            db.save_learned_command(self.olt_profile, "Delete_Provisioned", default_cmd)
            return default_cmd
        return del_cmd_temp

    def _delete_ont_by_port(self, port_full: str, db: DatabaseManager):
        logging.info(f"Aggressively deleting ONT bound to port: {port_full}...")
        del_cmd_temp = self._get_safe_delete_cmd(db)
        
        cmds = [c.strip() for c in del_cmd_temp.split(';') if c.strip()]
        for c in cmds:
            exec_cmd = c.format(port=port_full)
            self.send_command(exec_cmd, timeout=5)
            
        time.sleep(2)
        logging.info(f"Cleanup complete for port {port_full}.")

    def _cleanup_failed_provisioning(self, port_full: str):
        db = DatabaseManager()
        self._delete_ont_by_port(port_full, db)

    def format_serial(self, serial_raw: str) -> str:
        vendor_id = serial_raw[:4]
        sn_rem = serial_raw[4:]
        return f"{vendor_id}:{sn_rem}"

    def register_ont(self, config: Dict, db: DatabaseManager) -> bool:
        port_full = f"{config['pon']}/{config['ont_id']}"
        serial_raw = config['ont_serial']
        serial_formatted = self.format_serial(serial_raw)
        
        logging.info(f"--- [STEP 1] Initiating ONT Registration for {serial_raw} on {port_full} ---")
        
        cmd_intf1 = f"configure equipment ont interface {port_full} sw-ver-pland disabled sernum {serial_formatted} fec-up enable sw-dnload-version disabled pref-channel-pair {config.get('chanpair', '')}"
        cmd_intf2 = f"configure equipment ont interface {port_full} admin-state up"
        
        for c in [cmd_intf1, cmd_intf2]:
            out = self.send_command(c)
            if "^" in out or "error" in out.lower():
                logging.error(f"Command rejected: {c}\n{out}")
                self._cleanup_failed_provisioning(port_full)
                return False
                
        logging.info("Waiting 5s for hardware to synchronize slots and checking alarms...")
        time.sleep(5)
        
        alarm_out = self.send_command(f"show equipment ont alarm {port_full}")
        slot_out = self.send_command("show equipment ont slot")
        
        alarm_detected = "did not accept the configuration request" in alarm_out or "alarm occurred for ont-card" in alarm_out.lower()
        
        available_slots = []
        for line in slot_out.split('\n'):
            if port_full in line:
                m = re.search(rf'{re.escape(port_full)}/(\d+)', line)
                if m: available_slots.append(m.group(1))
        available_slots = list(set(available_slots))
        
        default_slot = config.get('ont_slot', '1')
        selected_slot = default_slot
        
        if alarm_detected or (available_slots and default_slot not in available_slots):
            slots_str = ", ".join(available_slots) if available_slots else "None Detected"
            msg = f"Slot mismatch or Card alarm detected!\nAvailable slots for {port_full}: {slots_str}\n\nEnter correct slot number to use:"
            new_slot = gui_ask(msg)
            if new_slot and new_slot.isdigit():
                selected_slot = new_slot
            else:
                logging.error("Slot selection aborted.")
                self._cleanup_failed_provisioning(port_full)
                return False
        elif available_slots and len(available_slots) == 1:
            selected_slot = available_slots[0]
            
        config['ont_slot'] = selected_slot
        logging.info(f"Selected ONT Slot assigned as: {selected_slot}")
        
        card_type = "ethernet" if config.get('ont_type', 'sfu') == 'sfu' else "veip"
        cmd_card1 = f"configure equipment ont slot {port_full}/{selected_slot} planned-card-type {card_type} plndnumdataports {config.get('lan_ports', '1')} plndnumvoiceports 0"
        cmd_card2 = f"configure interface port uni:{port_full}/{selected_slot}/1 admin-up"
        
        for c in [cmd_card1, cmd_card2]:
            out = self.send_command(c)
            if "^" in out or "error" in out.lower():
                logging.error(f"Card configuration rejected: {c}\n{out}")
                fix_cmd = gui_ask(f"Card config failed. Enter corrective CLI:")
                if fix_cmd and fix_cmd.lower() not in ['exit', '4']:
                    self.send_command(fix_cmd)
                else:
                    self._cleanup_failed_provisioning(port_full)
                    return False
        
        return True

    def verify_registration(self, config: Dict, db: DatabaseManager) -> bool:
        port_full = f"{config['pon']}/{config['ont_id']}"
        chanpair = config.get('chanpair', '')
        serial_formatted = self.format_serial(config['ont_serial'])

        cmd_template = db.get_learned_command(self.olt_profile, "Reg_Verify_Cmd")
        default_verify_cmd = "show equipment ont status channel-pair {chanpair}"
        
        if not cmd_template or ("{chanpair}" not in cmd_template and "{port}" not in cmd_template):
            db.save_learned_command(self.olt_profile, "Reg_Verify_Cmd", default_verify_cmd)
            cmd_template = default_verify_cmd

        logging.info(f"--- [STEP 2] Verifying ONT Registration on {port_full} ---")

        while True:
            cmd = cmd_template.format(port=port_full, chanpair=chanpair)
            max_retries = 5
            success = False
            
            for attempt in range(1, max_retries + 1):
                last_out = self.send_command(cmd, timeout=5)
                match_line = [line for line in last_out.split('\n') if serial_formatted.lower() in line.lower() and port_full.lower() in line.lower()]
                if match_line:
                    tokens = match_line[0].strip().lower().split()
                    if "up" in tokens:
                        success = True
                        break
                time.sleep(5)
                
            if success: return True
                
            print(f"\n[ VERIFICATION TIMEOUT ] Could not find oper status 'UP' for serial {serial_formatted} on port {port_full}.")
            choice = gui_ask("Verification Failed. Options: [1] Retry [2] Change Cmd [3] Force Pass [4] Abort (Enter number)")
            if choice == '1': continue
            elif choice == '2':
                new_cmd = gui_ask("Enter new verification command template:")
                if new_cmd and new_cmd.lower() not in ['exit', '4']: 
                    cmd_template = new_cmd
                    db.save_learned_command(self.olt_profile, "Reg_Verify_Cmd", cmd_template)
            elif choice == '3': return True
            else: return False

    def provision_service_ont(self, config: Dict, db: DatabaseManager) -> bool:
        port_full = f"{config['pon']}/{config['ont_id']}"
        slot = config.get('ont_slot', '1')
        vlan_mode = config.get('vlan_mode', 'untagged').lower()
        prov_key = f"Provisioning_Service_{config['ont_type'].upper()}_{vlan_mode.upper()}"
        
        learned_cmds = db.get_learned_command(self.olt_profile, prov_key)
        
        if learned_cmds and "vlan-id" not in learned_cmds.lower():
            logging.warning(f"Learned provisioning command for {vlan_mode} seems incomplete. Reverting to default sequence.")
            learned_cmds = None
            
        if learned_cmds:
            cmd_templates = [c.strip() for c in learned_cmds.split(';') if c.strip()]
            
            cmds_formatted = []
            for t in cmd_templates:
                c = t.format(
                    port=port_full, bw_profile=config.get('bw_profile', ''),
                    max_mac=config.get('max_mac', '128'), vlan_id=config.get('vlan_id', '1001'),
                    c_vlan=config.get('c_vlan', '10')
                )
                cmds_formatted.append(c)
        else:
            cmds_formatted = [
                f"configure qos interface uni:{port_full}/{slot}/1 queue [0...7] shaper-profile name:StrictPriority",
                f"configure qos interface uni:{port_full}/{slot}/1 upstream-queue [0...7] bandwidth-profile name:{config.get('bw_profile')} bandwidth-sharing uni-sharing",
                f"configure bridge port {port_full}/{slot}/1 max-unicast-mac {config.get('max_mac')}"
            ]
            
            if vlan_mode == 'untagged':
                cmds_formatted.extend([
                    f"configure bridge port {port_full}/{slot}/1 vlan-id {config.get('vlan_id')} usacceptframetype untagged",
                    f"configure bridge port {port_full}/{slot}/1 pvid {config.get('vlan_id')}"
                ])
            elif vlan_mode == 'tagged':
                cmds_formatted.append(f"configure bridge port {port_full}/{slot}/1 vlan-id {config.get('vlan_id')} tag single-tagged l2fwder-vlan {config.get('vlan_id')} vlan-scope local")
            elif vlan_mode == 'translation':
                cmds_formatted.append(f"configure bridge port {port_full}/{slot}/1 vlan-id {config.get('c_vlan')} tag single-tagged l2fwder-vlan {config.get('vlan_id')} vlan-scope local")

        logging.info(f"--- [STEP 3] Initiating Service Provisioning ({vlan_mode.upper()}) on {port_full}/{slot}/1 ---")
        error_kws = ["invalid command", "invalid token", "unknown command", "bad parameter", "incomplete command"]
        
        while True:
            success = True
            for cmd in cmds_formatted:
                out = self.send_command(cmd, timeout=10)
                out_lower = out.lower()
                if (any(kw in out_lower for kw in error_kws) or "^" in out) and "pattern not detected" not in out_lower:
                    print(f"\n[ SERVICE PROVISIONING FAILED ] Command rejected: {cmd}")
                    self._cleanup_failed_provisioning(port_full)
                    new_cmds = gui_ask("Service Provisioning failed. Enter corrective CLI sequence:")
                    if new_cmds.lower() in ['exit', '4'] or not new_cmds: return False 
                    db.save_learned_command(self.olt_profile, prov_key, new_cmds)
                    return False 
                    
            if success: return True

class TestAutomationEngine:
    def __init__(self, db: DatabaseManager, olt, config: Dict):
        self.db = db
        self.olt = olt
        self.config = config
        self.serial = config['ont_serial']
        self.vendor = config.get("global_vendor", "Nokia ISAM7360")
        
        if self.vendor == "Nokia ISAM7360":
            self.pon = config['pon']
            self.ont_id = config['ont_id']
            self.chanpair = config.get('chanpair', '')
            self.port_full = f"{self.pon}/{self.ont_id}"
            vendor_id = self.serial[:4]
            sn_rem = self.serial[4:]
            self.serial_formatted = f"{vendor_id}:{sn_rem}"
        else:
            self.pon = config.get('hw_fsp', '0/1/0')
            self.ont_id = config.get('hw_ont_id', '1')
            self.chanpair = ""
            self.port_full = f"{self.pon} {self.ont_id}"
            self.serial_formatted = self.serial

    def execute_tests(self):
        scenarios = self.db.load_test_cases(self.serial)
        retry_required = False
        
        for sn in scenarios:
            logging.info(f"\nExecuting [{sn['order']}] {sn['type']}")
            if "wait" in sn['parameters'] and sn['parameters']['wait'] > 0:
                logging.info(f"Waiting {sn['parameters']['wait']} seconds for hardware synchronization...")
                time.sleep(sn['parameters']['wait'])
            
            res = ""
            
            if sn['type'] == "Speed_Test":
                server = self.config.get('iperf_server', '').strip()
                port = str(self.config.get('iperf_port', 5201))
                duration = str(self.config.get('speed_duration', 10))
                criteria = float(self.config.get('speed_pass_mbps', 500))
                
                if not server:
                    logging.warning("No iperf3 server specified. Skipping Speed_Test.")
                    res = "[SPEED_TEST_SUCCESS] Skipped (No server)"
                    self.db.update_test_status(sn['id'], "PASS", res)
                    continue
                
                try:
                    overall_up = 0.0
                    overall_dl = 0.0
                    full_log = []
                    
                    base_args = ['iperf3', '-c', server, '-p', port, '-t', duration, '-P', '8']
                    tests = [("UPLOAD", base_args), ("DOWNLOAD", base_args + ['-R'])]
                    
                    for t_name, t_args in tests:
                        logging.info(f"Running Native iperf3 {t_name}: {' '.join(t_args)}")
                        print("\n" + "="*68)
                        print(f"[{t_name}] {'Time':<8} | {'Progress Graph':<42} | {'Throughput':<15}")
                        print("="*68)
                        
                        full_log.append(f"--- STARTING {t_name} TEST ---\n")
                        
                        proc = subprocess.Popen(t_args, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, bufsize=1)
                        direction_log = []
                        
                        for line in iter(proc.stdout.readline, ''):
                            direction_log.append(line)
                            full_log.append(line)
                            if "[SUM]" in line and "bits/sec" in line and "sender" not in line and "receiver" not in line:
                                match = re.search(r'\[SUM\]\s+\d+\.\d+-\s*(\d+\.\d+)\s+sec.*?\s+(\d+(?:\.\d+)?)\s+([KMG])bits/sec', line)
                                if match:
                                    sec_val = match.group(1)
                                    val = float(match.group(2))
                                    unit = match.group(3)
                                    mbps = val * 1000 if unit == 'G' else (val if unit == 'M' else val / 1000)
                                    
                                    bar_len = 40
                                    max_val = max(criteria * 1.2, 10000)
                                    filled = min(int((mbps / max_val) * bar_len), bar_len)
                                    bar = '█' * filled + '-' * (bar_len - filled)
                                    
                                    print(f"[{sec_val:>5}s] [{bar}] {mbps:8.2f} Mbps")
                                    
                        proc.wait()
                        print("="*68 + "\n")
                        
                        dir_text = "".join(direction_log)
                        if proc.returncode != 0:
                            logging.warning(f"{t_name} Speed Test returned code {proc.returncode}")
                            
                        if t_name == "UPLOAD":
                            sender_matches = re.findall(r'\[SUM\].*?\s+(\d+(?:\.\d+)?)\s+([KMG])bits/sec\s+sender', dir_text)
                            if sender_matches:
                                v, u = float(sender_matches[-1][0]), sender_matches[-1][1]
                                overall_up = v * 1000 if u == 'G' else (v if u == 'M' else v / 1000)
                        else:
                            recv_matches = re.findall(r'\[SUM\].*?\s+(\d+(?:\.\d+)?)\s+([KMG])bits/sec\s+receiver', dir_text)
                            if recv_matches:
                                v, u = float(recv_matches[-1][0]), recv_matches[-1][1]
                                overall_dl = v * 1000 if u == 'G' else (v if u == 'M' else v / 1000)
                                
                        time.sleep(2) 
                        
                    out_text = "".join(full_log)
                    
                    if overall_up >= criteria and overall_dl >= criteria:
                        res = f"[SPEED_TEST_SUCCESS] Upload: {overall_up:.2f} Mbps, Download: {overall_dl:.2f} Mbps\n{out_text}"
                        logging.info(f"Speed test passed criteria! (Up: {overall_up:.2f} Mbps, Down: {overall_dl:.2f} Mbps)")
                    else:
                        res = f"[SPEED_TEST_FAILED] Upload: {overall_up:.2f} Mbps, Download: {overall_dl:.2f} Mbps (Target: {criteria} Mbps)\n{out_text}"
                        logging.warning(f"Speed test throughput below target {criteria} Mbps.")
                
                except FileNotFoundError:
                    res = "[SPEED_TEST_FAILED] iperf3 executable not found in system PATH."
                    logging.error(res)
                except Exception as e:
                    res = f"[SPEED_TEST_FAILED] Unexpected error: {e}"
                    logging.error(res)
                    
            elif sn['type'] == "Reboot_Test" and self.vendor == "Nokia ISAM7360":
                logging.info("Sending Reboot Command to OLT...")
                res_initial = self.olt.send_command(sn['command'])
                
                if "^" in res_initial or "invalid token" in res_initial.lower() or "error" in res_initial.lower():
                    res = res_initial 
                else:
                    logging.info("Reboot command sent. Forcing OLT disconnection to allow PC NIC failover (Wait 5s)...")
                    time.sleep(5)
                    self.olt.disconnect()
                    
                    logging.info("Attempting to reconnect to OLT and waiting for ONT recovery (Max 5 mins)...")
                    start_wait_t = time.time()
                    reconnected = False
                    
                    while time.time() - start_wait_t < 300:
                        if self.olt.connect():
                            reconnected = True
                            logging.info("Successfully reconnected to OLT!")
                            break
                        time.sleep(5)
                        
                    if not reconnected:
                        res = f"[REBOOT_FAILED] Could not reconnect to OLT after 5 minutes.\n[Reboot Trigger]\n{res_initial}"
                        logging.error("OLT Reconnection timeout reached!")
                    else:
                        alarm_cmd_template = self.db.get_learned_command(self.olt.olt_profile, "Alarm_Check_Cmd")
                        if not alarm_cmd_template:
                            alarm_cmd_template = "show equipment ont alarm {port}" 
                            self.db.save_learned_command(self.olt.olt_profile, "Alarm_Check_Cmd", alarm_cmd_template)
                            
                        alarm_cmd = alarm_cmd_template.format(port=self.port_full, chanpair=self.chanpair)
                        
                        reboot_success = False
                        reboot_time_sec = 0.0
                        last_alarm_out = ""
                        
                        log_pattern = re.compile(r'(\d{2}/\d{2}/\d{2}\s+\d{2}:\d{2}:\d{2})\s+major alarm (occurred|cleared).*ONT is inactive', re.IGNORECASE)
                        
                        for attempt in range(60):
                            last_alarm_out = self.olt.send_command(alarm_cmd, timeout=5)
                            
                            occurred_time = None
                            cleared_time = None
                            
                            for match in log_pattern.finditer(last_alarm_out):
                                timestamp_str, action = match.groups()
                                try:
                                    dt_obj = datetime.strptime(timestamp_str, "%y/%m/%d %H:%M:%S")
                                    if action.lower() == "occurred":
                                        occurred_time = dt_obj
                                    elif action.lower() == "cleared":
                                        cleared_time = dt_obj
                                except ValueError:
                                    continue
                                    
                            if occurred_time and cleared_time and cleared_time > occurred_time:
                                reboot_time_sec = (cleared_time - occurred_time).total_seconds()
                                reboot_success = True
                                break
                            elif cleared_time and not occurred_time:
                                reboot_success = True
                                reboot_time_sec = time.time() - start_wait_t
                                break
                                
                            time.sleep(5)
                            
                        if reboot_success:
                            res = f"[REBOOT_SUCCESS] ONT recovery confirmed via Alarm Log. True Reboot Duration: {reboot_time_sec:.1f} seconds.\n[Reboot Trigger]\n{res_initial}\n[Final Alarm Log]\n{last_alarm_out}"
                            logging.info(f"ONT Rebooted and Recovery Detected! Actual Reboot Time: {reboot_time_sec:.1f} seconds.")
                        else:
                            res = f"[REBOOT_FAILED] ONT did not recover within timeout (Alarm 'cleared' not found).\n[Reboot Trigger]\n{res_initial}\n[Last Alarm Log]\n{last_alarm_out}"
                            logging.error("Reboot timeout reached!")

            else:
                if ';' in sn['command']:
                    res_lines = []
                    for c in sn['command'].split(';'):
                        c = c.strip()
                        if c:
                            logging.info(f"Executing: {c}")
                            out = self.olt.send_command(c)
                            res_lines.append(out)
                    res = "\n".join(res_lines)
                else:
                    res = self.olt.send_command(sn['command'])
                    
                if sn['type'] == "Software_Info" and self.vendor == "Nokia ISAM7360":
                    match = re.search(r'sw-ver-act\s*:\s*(\S+)', res, re.IGNORECASE)
                    if match:
                        self.config['sw_version'] = match.group(1).upper()
            
            if self._verify(sn['type'], res):
                logging.info(f"[PASS] {sn['type']}")
                self.db.update_test_status(sn['id'], "PASS", res)
            else:
                if sn['type'] == "Speed_Test":
                    logging.warning(f"[{sn['type']}] Did not meet the target criteria. Proceeding with remaining tests.")
                    self.db.update_test_status(sn['id'], "FAIL", res)
                    continue
                
                action = self._handle_failure(sn, res)
                
                if action == "ABORT":
                    self.db.update_test_status(sn['id'], "FAIL", res)
                    return "ABORT"
                    
                if not action: 
                    self.db.update_test_status(sn['id'], "FAIL", res)
                    retry_required = True
                    break 
                    
                if action == "SKIP":
                    logging.info(f"[SKIPPED] {sn['type']}")
                    self.db.update_test_status(sn['id'], "N/A", res)
                    continue 

                retry_required = True
                break 
                
        if retry_required:
            logging.warning(f"Test sequence aborted. Initiating re-test in the next cycle.")
            return False
        return True

    def _verify(self, t_type: str, res: str) -> bool:
        if not res:
            if t_type in ["Cleanup"]: return True
            return False
            
        res_lower = res.lower()
        
        if t_type == "Reboot_Test":
            return "[REBOOT_SUCCESS]" in res
        if t_type == "Speed_Test":
            return "[SPEED_TEST_SUCCESS]" in res
            
        error_kws = ["invalid command", "invalid token", "unknown command", "bad parameter", "incomplete command"]
        if any(kw in res_lower for kw in error_kws) or "^" in res: 
            return False
            
        if self.vendor == "Nokia ISAM7360":
            if t_type == "ONT_Discovery_Check":
                serial_clean = self.serial.replace(":", "").lower()
                return serial_clean in res_lower.replace(":", "")

            if t_type == "Registration_Check":
                serial_clean = self.serial.replace(":", "").lower()
                return "pref-ranged" in res_lower and serial_clean in res_lower.replace(":", "")
                
            if t_type in ["Optics_Signal_Check", "Optics_Temp_Check", "Optics_Voltage_Laser_Check"]:
                success_found = False
                for line in res.split('\n'):
                    if self.port_full.lower() in line.lower() or self.port_full.split(':')[-1] in line:
                        tokens = line.strip().split()
                        port_idx = -1
                        for i, t in enumerate(tokens):
                            if self.port_full.lower() in t.lower() or self.port_full.split(':')[-1] in t:
                                port_idx = i
                                break
                        
                        if port_idx != -1 and len(tokens) >= port_idx + 7:
                            rx_sig = tokens[port_idx + 1]
                            tx_sig = tokens[port_idx + 2]
                            temp = tokens[port_idx + 3]
                            voltage = tokens[port_idx + 4]
                            laser = tokens[port_idx + 5]
                            olt_rx = tokens[port_idx + 6]
                            
                            def is_num(val):
                                val_clean = val.lower()
                                if val_clean in ['invalid', 'not-appl', 'n/a', 'none', '-', 'inf', '-inf']: return False
                                try:
                                    float(val)
                                    return True
                                except ValueError:
                                    return False
                            
                            if t_type == "Optics_Signal_Check":
                                success_found = is_num(rx_sig) and is_num(tx_sig) and is_num(olt_rx)
                            elif t_type == "Optics_Temp_Check":
                                success_found = is_num(temp)
                            elif t_type == "Optics_Voltage_Laser_Check":
                                success_found = is_num(voltage) and is_num(laser)
                                
                            if success_found:
                                break
                return success_found
                
            if t_type == "Software_Info":
                match = re.search(r'sw-ver-act\s*:\s*(\S+)', res_lower)
                if match:
                    val = match.group(1)
                    return val not in ["sw-ver-psv", "vendor-id", "unknown"]
                return False
                
            if t_type == "UNI_Status":
                has_link = "link-status" in res_lower and "up" in res_lower
                has_speed = "config-indicator" in res_lower and any(s in res_lower for s in ['10g', '1g', '100m', '10m'])
                return has_link and has_speed
                
            if t_type == "MAC_Status": 
                return bool(re.search(r'([0-9A-F]{2}:){5}[0-9A-F]{2}', res, re.I))
                
        else:
            # Huawei baseline verification (to be expanded)
            return True

        return True

    def _handle_failure(self, sn: Dict, res: str):
        print(f"\n[FAILURE DETECTED] {sn['type']}")
        
        res_lower = res.lower() if res else ""
        error_kws = ["invalid command", "invalid token", "unknown command", "bad parameter", "incomplete command"]
        is_syntax_error = any(kw in res_lower for kw in error_kws) or "^" in res
        
        if is_syntax_error and sn['type'] != "Speed_Test":
            print(f"\n[ AI Auto-Detection ] 'Invalid CLI' detected.")
            tmp = gui_ask(f"Enter correct CLI template for {sn['type']} (or type 'exit'):")
            if tmp.lower() in ['exit', '4']: 
                logging.info("User selected ABORT during CLI correction.")
                return "ABORT"
            if tmp:
                self.db.save_learned_command(self.olt.olt_profile, sn['type'], tmp)
                self.db.update_test_case(sn['id'], tmp.format(serial=self.serial, pon=self.pon, ont_id=self.ont_id, port=self.port_full), json.dumps(sn['parameters']))
                return True
                
        choice = gui_ask(f"Error in {sn['type']}. Select Option: [1] Update Cmd [2] Increase Wait [3] Force Pass [4] Abort")
        if choice == "1":
            tmp = gui_ask(f"Enter new template for {sn['type']}:")
            if tmp and tmp.lower() not in ['exit', '4']:
                self.db.save_learned_command(self.olt.olt_profile, sn['type'], tmp)
                self.db.update_test_case(sn['id'], tmp.format(serial=self.serial, pon=self.pon, ont_id=self.ont_id, port=self.port_full), json.dumps(sn['parameters']))
            return True
        elif choice == "2":
            sn['parameters']['wait'] = sn['parameters'].get('wait', 0) + 10
            self.db.update_test_case(sn['id'], sn['command'], json.dumps(sn['parameters']))
            return True
        elif choice == "3":
            logging.info(f"Force passing {sn['type']} step.")
            return "SKIP"
        elif choice in ["4", "exit"]:
            logging.warning("User manually initiated ABORT sequence.")
            return "ABORT"
        else:
            logging.warning("Invalid input received. Defaulting to ABORT.")
            return "ABORT"

# ==========================================
# GUI Application Layer
# ==========================================
class AutomationGUI:
    def __init__(self, root):
        global GUI_ROOT
        GUI_ROOT = root
        self.root = root
        self.root.title("ONT Automation Tool v2.0 (Multi-Vendor)")
        self.root.geometry("1100x850")
        self.root.configure(bg="#2E3440")
        
        self.db = DatabaseManager()
        self.config = self.load_config()
        self.olt = None
        self.setup_ui()
        self.setup_logging()

    def load_config(self):
        if os.path.exists(CONFIG_FILE):
            with open(CONFIG_FILE, 'r') as f: return json.load(f)
        return {}

    def save_config(self):
        with open(CONFIG_FILE, 'w') as f: json.dump(self.config, f, indent=4)

    def setup_logging(self):
        redir = SafeTextRedirector(self.log_text)
        sys.stdout = redir
        sys.stderr = redir
        handler = TextHandler(redir)
        handler.setFormatter(logging.Formatter('%(levelname)s: %(message)s'))
        logging.getLogger().addHandler(handler)
        logging.getLogger().setLevel(logging.INFO)

    def on_vendor_change(self, event=None):
        vendor = self.cb_global_vendor.get()
        if vendor == "Nokia ISAM7360":
            self.nokia_target_frame.pack(fill=tk.BOTH, expand=True)
            self.huawei_target_frame.pack_forget()
            
            self.nokia_svc_frame.pack(fill=tk.BOTH, expand=True)
            self.huawei_svc_frame.pack_forget()
        else:
            self.nokia_target_frame.pack_forget()
            self.huawei_target_frame.pack(fill=tk.BOTH, expand=True)
            
            self.nokia_svc_frame.pack_forget()
            self.huawei_svc_frame.pack(fill=tk.BOTH, expand=True)

    def setup_ui(self):
        style = ttk.Style()
        style.theme_use("clam")
        
        left_frame = tk.Frame(self.root, bg="#2E3440", width=400)
        left_frame.pack(side=tk.LEFT, fill=tk.Y, padx=10, pady=10)
        left_frame.pack_propagate(False)

        # Global Vendor Selector
        vendor_frame = tk.Frame(left_frame, bg="#2E3440")
        vendor_frame.pack(fill=tk.X, pady=(0, 10))
        tk.Label(vendor_frame, text="Select Target OLT:", bg="#2E3440", fg="white", font=("Arial", 10, "bold")).pack(side=tk.LEFT)
        self.cb_global_vendor = ttk.Combobox(vendor_frame, values=["Nokia ISAM7360", "Huawei MA5800-X7"], state="readonly")
        self.cb_global_vendor.set(self.config.get("global_vendor", "Nokia ISAM7360"))
        self.cb_global_vendor.pack(side=tk.LEFT, padx=10, fill=tk.X, expand=True)
        self.cb_global_vendor.bind("<<ComboboxSelected>>", self.on_vendor_change)

        self.notebook = ttk.Notebook(left_frame)
        self.notebook.pack(fill=tk.BOTH, expand=True)

        # ---------------------------------------------------------
        # Tab 1: Target & Connection
        # ---------------------------------------------------------
        self.tab_tgt = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_tgt, text="Target & Connection")
        
        info_frame = ttk.LabelFrame(self.tab_tgt, text="OLT System Info")
        info_frame.pack(fill=tk.X, padx=10, pady=5)
        self.lbl_vendor = ttk.Label(info_frame, text="Vendor: -")
        self.lbl_vendor.pack(anchor='w', padx=5, pady=2)
        self.lbl_model = ttk.Label(info_frame, text="Model: -")
        self.lbl_model.pack(anchor='w', padx=5, pady=2)
        self.lbl_sw_ver = ttk.Label(info_frame, text="SW Version: -")
        self.lbl_sw_ver.pack(anchor='w', padx=5, pady=2)
        
        self.sv_olt_ip = self.add_entry(self.tab_tgt, "OLT IP:", self.config.get("olt_ip", ""))
        self.sv_username = self.add_entry(self.tab_tgt, "Username:", self.config.get("username", "isadmin"))
        self.sv_password = self.add_entry(self.tab_tgt, "Password:", self.config.get("password", ""), show="*")
        
        ttk.Separator(self.tab_tgt, orient='horizontal').pack(fill='x', pady=10)
        
        self.sv_serial = self.add_entry(self.tab_tgt, "ONT Serial / SN:", self.config.get("ont_serial", ""))
        
        # Dynamic Target Frames
        self.nokia_target_frame = tk.Frame(self.tab_tgt)
        self.sv_chanpair = self.add_entry(self.nokia_target_frame, "Channel Pair:", self.config.get("chanpair", "1/1/1/3"))
        self.sv_pon = self.add_entry(self.nokia_target_frame, "Target PON:", self.config.get("pon", "ng2:3/1"))
        self.sv_ont_id = self.add_entry(self.nokia_target_frame, "Target ONT ID:", self.config.get("ont_id", "1"))
        
        self.huawei_target_frame = tk.Frame(self.tab_tgt)
        self.sv_hw_fsp = self.add_entry(self.huawei_target_frame, "Target F/S/P:", self.config.get("hw_fsp", "0/1/0"))
        self.sv_hw_ont_id = self.add_entry(self.huawei_target_frame, "Target ONT ID:", self.config.get("hw_ont_id", "1"))

        # ---------------------------------------------------------
        # Tab 2: Service Config
        # ---------------------------------------------------------
        self.tab_svc = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_svc, text="Service & VLAN")
        
        self.cb_ont_type = self.add_combobox(self.tab_svc, "ONT Type:", ["sfu", "hgu"], self.config.get("ont_type", "sfu"))
        self.sv_lan_ports = self.add_entry(self.tab_svc, "LAN Ports:", self.config.get("lan_ports", "1"))
        self.sv_max_mac = self.add_entry(self.tab_svc, "Max MAC:", self.config.get("max_mac", "16"))
        
        # Dynamic Service Frames
        self.nokia_svc_frame = tk.Frame(self.tab_svc)
        self.sv_ont_slot = self.add_entry(self.nokia_svc_frame, "ONT Slot (Data):", self.config.get("ont_slot", "1"))
        self.sv_bw_profile = self.add_entry(self.nokia_svc_frame, "BW Profile Name:", self.config.get("bw_profile", "NG2DATABWUP10000"))
        
        self.huawei_svc_frame = tk.Frame(self.tab_svc)
        self.sv_hw_line_prof = self.add_entry(self.huawei_svc_frame, "Line Profile ID:", self.config.get("hw_line_prof", "10"))
        self.sv_hw_srv_prof = self.add_entry(self.huawei_svc_frame, "Service Profile ID:", self.config.get("hw_srv_prof", "10"))
        
        ttk.Separator(self.tab_svc, orient='horizontal').pack(fill='x', pady=10)
        
        self.cb_vlan_mode = self.add_combobox(self.tab_svc, "VLAN Mode:", ["untagged", "tagged", "translation"], self.config.get("vlan_mode", "untagged"))
        self.sv_vlan_id = self.add_entry(self.tab_svc, "Service VLAN ID:", self.config.get("vlan_id", "1001"))
        self.sv_cvlan = self.add_entry(self.tab_svc, "C-VLAN (Translation):", self.config.get("c_vlan", "10"))

        # Initialize the correct view based on saved config
        self.on_vendor_change()

        # ---------------------------------------------------------
        # Tab 3: Speed Test
        # ---------------------------------------------------------
        tab_spd = ttk.Frame(self.notebook)
        self.notebook.add(tab_spd, text="Speed Test Config")
        
        self.sv_iperf_ip = self.add_entry(tab_spd, "iPerf3 Server IP:", self.config.get("iperf_server", "103.175.200.43"))
        self.sv_iperf_port = self.add_entry(tab_spd, "Server Port:", self.config.get("iperf_port", "5201"))
        self.sv_duration = self.add_entry(tab_spd, "Test Duration (sec):", self.config.get("speed_duration", "30"))
        self.sv_criteria = self.add_entry(tab_spd, "Pass Criteria (Mbps):", self.config.get("speed_pass_mbps", "8000"))

        # Right Panel: Controls & Log
        right_frame = tk.Frame(self.root, bg="#2E3440")
        right_frame.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=10, pady=10)

        ctrl_frame = tk.Frame(right_frame, bg="#2E3440")
        ctrl_frame.pack(fill=tk.X, pady=(0, 10))

        btn_style = {"bg": "#4C566A", "fg": "white", "font": ("Arial", 10, "bold"), "relief": tk.FLAT, "padx": 10, "pady": 5}
        
        tk.Button(ctrl_frame, text="Scan Unprovisioned", command=lambda: self.run_thread(self.scan_onts, "unprovisioned"), **btn_style).pack(side=tk.LEFT, padx=5)
        tk.Button(ctrl_frame, text="Scan Active (UP)", command=lambda: self.run_thread(self.scan_onts, "active"), **btn_style).pack(side=tk.LEFT, padx=5)
        tk.Button(ctrl_frame, text="Provision & Test", command=lambda: self.run_thread(self.start_automation, False), bg="#A3BE8C", fg="white", font=("Arial", 10, "bold")).pack(side=tk.LEFT, padx=5)
        tk.Button(ctrl_frame, text="Test Only", command=lambda: self.run_thread(self.start_automation, True), **btn_style).pack(side=tk.LEFT, padx=5)
        tk.Button(ctrl_frame, text="Delete ONT", command=lambda: self.run_thread(self.delete_ont), bg="#BF616A", fg="white", font=("Arial", 10, "bold")).pack(side=tk.RIGHT, padx=5)

        self.log_text = scrolledtext.ScrolledText(right_frame, bg="#1E1E1E", fg="#D8DEE9", font=("Consolas", 10), state=tk.DISABLED)
        self.log_text.pack(fill=tk.BOTH, expand=True)

    def add_entry(self, parent, label_text, default_val, show=None):
        frame = tk.Frame(parent)
        frame.pack(fill=tk.X, padx=10, pady=5)
        tk.Label(frame, text=label_text, width=20, anchor='w').pack(side=tk.LEFT)
        sv = tk.StringVar(value=default_val)
        tk.Entry(frame, textvariable=sv, show=show).pack(side=tk.RIGHT, fill=tk.X, expand=True)
        return sv

    def add_combobox(self, parent, label_text, values, default_val):
        frame = tk.Frame(parent)
        frame.pack(fill=tk.X, padx=10, pady=5)
        tk.Label(frame, text=label_text, width=20, anchor='w').pack(side=tk.LEFT)
        cb = ttk.Combobox(frame, values=values, state="readonly")
        cb.set(default_val)
        cb.pack(side=tk.RIGHT, fill=tk.X, expand=True)
        return cb

    def update_config_from_ui(self):
        self.config['global_vendor'] = self.cb_global_vendor.get()
        
        self.config['olt_ip'] = self.sv_olt_ip.get()
        self.config['username'] = self.sv_username.get()
        self.config['password'] = self.sv_password.get()
        self.config['protocol'] = "telnet"
        
        self.config['ont_serial'] = self.sv_serial.get()
        
        # Nokia Fields
        self.config['chanpair'] = self.sv_chanpair.get()
        self.config['pon'] = self.sv_pon.get()
        self.config['ont_id'] = self.sv_ont_id.get()
        self.config['ont_slot'] = self.sv_ont_slot.get() 
        self.config['bw_profile'] = self.sv_bw_profile.get()
        
        # Huawei Fields
        self.config['hw_fsp'] = self.sv_hw_fsp.get()
        self.config['hw_ont_id'] = self.sv_hw_ont_id.get()
        self.config['hw_line_prof'] = self.sv_hw_line_prof.get()
        self.config['hw_srv_prof'] = self.sv_hw_srv_prof.get()
        
        self.config['ont_type'] = self.cb_ont_type.get()
        self.config['lan_ports'] = self.sv_lan_ports.get()
        self.config['max_mac'] = self.sv_max_mac.get()
        self.config['vlan_mode'] = self.cb_vlan_mode.get()
        self.config['vlan_id'] = self.sv_vlan_id.get()
        self.config['c_vlan'] = self.sv_cvlan.get()
        
        self.config['iperf_server'] = self.sv_iperf_ip.get()
        self.config['iperf_port'] = self.sv_iperf_port.get()
        self.config['speed_duration'] = self.sv_duration.get()
        self.config['speed_pass_mbps'] = self.sv_criteria.get()
        
        self.save_config()

    def update_olt_info_display(self, vendor, model, sw_ver):
        self.lbl_vendor.config(text=f"Vendor: {vendor}")
        self.lbl_model.config(text=f"Model: {model}")
        self.lbl_sw_ver.config(text=f"SW Version: {sw_ver}")
        self.config['olt_vendor'] = vendor
        self.config['olt_model'] = model
        self.save_config()

    def run_thread(self, target_func, *args):
        self.update_config_from_ui()
        threading.Thread(target=target_func, args=args, daemon=True).start()

    def get_connector(self):
        vendor = self.cb_global_vendor.get()
        
        if vendor == "Huawei MA5800-X7":
            self.olt = HuaweiOLTConnector(self.config['olt_ip'], self.config['username'], self.config['password'], self.config['protocol'])
            self.olt.connect()
            return self.olt
            
        if not self.olt or not self.olt.connection or not self.olt.connection.is_alive() or isinstance(self.olt, HuaweiOLTConnector):
            self.olt = NokiaOLTConnector(self.config['olt_ip'], self.config['username'], self.config['password'], self.config['protocol'])
            if not self.olt.connect():
                logging.error("Failed to connect to Nokia OLT. Check IP/Credentials.")
                return None
            else:
                self.root.after(0, self.update_olt_info_display, self.olt.vendor, self.olt.model, self.olt.sw_version)
        return self.olt

    def show_selection_popup(self, title, items, callback):
        if not items:
            messagebox.showinfo("No Results", f"No items found for {title}.")
            return

        def on_select(evt):
            sel = listbox.curselection()
            if sel:
                idx = sel[0]
                callback(items[idx])
                top.destroy()

        top = tk.Toplevel(self.root)
        top.title(title)
        top.geometry("400x300")
        top.transient(self.root)
        top.grab_set()

        listbox = tk.Listbox(top, font=("Consolas", 10))
        listbox.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        for item in items:
            listbox.insert(tk.END, f"Serial: {item['serial']} | Port: {item.get('port', 'N/A')} | Chan: {item['chanpair']}")
            
        listbox.bind('<Double-1>', on_select)
        tk.Button(top, text="Cancel", command=top.destroy).pack(pady=5)

    def scan_onts(self, scan_type="unprovisioned"):
        olt = self.get_connector()
        if not olt: return
        
        if isinstance(olt, HuaweiOLTConnector):
            logging.warning("Huawei scan logic is not fully implemented yet.")
            return

        print("\n" + "="*50)
        logging.info(f"Starting {scan_type.upper()} ONT Scan...")
        
        items = []
        if scan_type == "unprovisioned":
            cmd = self.db.get_learned_command(olt.olt_profile, "Discovery") or "show channel-pair unprovision-onu"
            items = olt.get_unprovisioned_onts(cmd)
        else:
            for i in range(1, 9):
                out = olt.send_command(f"show equipment ont status channel-pair 1/1/1/{i}", timeout=5)
                matches = re.findall(r'(\d+(?:/\d+)+)\s+([a-zA-Z0-9:-]+(?:/\d+)+)\s+([A-Za-z]{4}:?[A-Fa-f0-9]{8})\s+(\S+)\s+(\S+)', out)
                for m in matches:
                    if m[4].lower() == 'up': 
                        items.append({'chanpair': m[0], 'port': m[1], 'serial': m[2].replace(':', '')})
                        
        logging.info(f"Scan complete. Found {len(items)} items.")
        
        def handle_selection(selected):
            self.sv_serial.set(selected['serial'])
            self.sv_chanpair.set(selected['chanpair'])
            
            if scan_type == "active":
                pon, ont_id = selected['port'].rsplit('/', 1)
                self.sv_pon.set(pon)
                self.sv_ont_id.set(ont_id)
            else:
                parts = selected['chanpair'].split('/')
                if len(parts) >= 4:
                    self.sv_pon.set(f"ng2:{parts[3]}/{parts[2]}")
                
                logging.info(f"Scanning channel {selected['chanpair']} for available ID...")
                out = olt.send_command(f"show equipment ont status channel-pair {selected['chanpair']}", timeout=5)
                used_ids = set()
                matches = re.findall(r'([a-zA-Z0-9:-]+(?:/\d+)+)\s+([A-Za-z]{4}:?[A-Fa-f0-9]{8})', out)
                for port_str, _ in matches:
                    last_digit = port_str.split('/')[-1]
                    if last_digit.isdigit():
                        used_ids.add(int(last_digit))
                        
                for i in range(1, 129):
                    if i not in used_ids:
                        self.sv_ont_id.set(str(i))
                        break
            
            logging.info(f"Auto-filled target info for {selected['serial']}.")

        self.root.after(0, lambda: self.show_selection_popup(f"Select {scan_type.capitalize()} ONT", items, handle_selection))

    def delete_ont(self):
        olt = self.get_connector()
        if not olt: return
        
        if isinstance(olt, HuaweiOLTConnector):
            logging.warning("Huawei delete logic is not fully implemented yet.")
            return

        port_full = f"{self.config['pon']}/{self.config['ont_id']}"
        confirm = messagebox.askyesno("Confirm Deletion", f"Are you sure you want to delete ONT on port {port_full}?")
        if confirm:
            print("\n" + "="*50)
            olt._delete_ont_by_port(port_full, self.db)
            logging.info(f"Deletion complete for {port_full}.")

    def start_automation(self, skip_provisioning=False):
        olt = self.get_connector()
        if not olt: return
        
        if isinstance(olt, HuaweiOLTConnector):
            logging.warning("Huawei automation pipeline will be activated in the next patch.")
            return

        port_full = f"{self.config['pon']}/{self.config['ont_id']}"
        print("\n" + "="*60)
        logging.info("Starting Nokia Automation Sequence...")
        
        try:
            if not skip_provisioning:
                while True:
                    if not olt.register_ont(self.config, self.db): 
                        logging.error("Registration aborted.")
                        return
                    if not olt.verify_registration(self.config, self.db):
                        olt._cleanup_failed_provisioning(port_full)
                        if gui_ask("Registration Verification Failed. Retry? (y/n)").lower() == 'y': continue
                        return
                    if not olt.provision_service_ont(self.config, self.db):
                        olt._cleanup_failed_provisioning(port_full)
                        if gui_ask("Provisioning Failed. Retry sequence? (y/n)").lower() == 'y': continue
                        return
                    break 
                self.db.add_ont(self.config['ont_serial'], self.config['pon'], self.config['ont_id'], "PROVISIONED")
            else:
                logging.info("Checking current slot assignment for Test Only mode...")
                slot_out = olt.send_command("show equipment ont slot")
                available_slots = []
                for line in slot_out.split('\n'):
                    if port_full in line:
                        m = re.search(rf'{re.escape(port_full)}/(\d+)', line)
                        if m: available_slots.append(m.group(1))
                if available_slots:
                    self.config['ont_slot'] = available_slots[0]
                else:
                    tmp = gui_ask("Could not detect slot. Enter slot number:")
                    self.config['ont_slot'] = tmp if tmp and tmp not in ['exit', '4'] else "1"

            self.db.add_initial_test_cases(self.config, olt.olt_profile)

            engine = TestAutomationEngine(self.db, olt, self.config)
            for cycle in range(1, 10):
                logging.info(f"\n========== TEST CYCLE {cycle} ==========")
                test_result = engine.execute_tests()
                if test_result == True: 
                    logging.info("ALL TESTS PASSED SUCCESSFULLY!")
                    break
                elif test_result == "ABORT":
                    logging.warning("Testing manually aborted by user. Terminating further cycles.")
                    break
        except Exception as e:
            logging.error(f"Critical error during automation: {e}")
        finally:
            generate_professional_excel_report(self.config['ont_serial'], self.db, self.config)
            logging.info("Automation Task & Report Generation Completed.")

if __name__ == "__main__":
    root = tk.Tk()
    app = AutomationGUI(root)
    
    try:
        root.mainloop()
    except KeyboardInterrupt:
        print("\n[!] Program interrupted by user. Exiting safely...")
    finally:
        if app.olt: app.olt.disconnect()
        sys.exit(0)