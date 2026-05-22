import logging
import re
from datetime import datetime
from typing import Dict

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.marker import Marker

    EXCEL_SUPPORT = True
except ImportError:
    logging.warning(
        "openpyxl is not installed! Excel report generation will be skipped. Run 'pip install openpyxl'"
    )
    EXCEL_SUPPORT = False


def generate_professional_excel_report(serial: str, db, config: Dict):
    """Generates a highly formatted 3-sheet Excel report with live speed graphs."""
    if not EXCEL_SUPPORT:
        return

    cases = db.load_test_cases(serial)
    if not cases:
        return

    filename = f"ONT_Test_Result_{serial}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    try:
        wb = openpyxl.Workbook()

        bold_font = Font(bold=True)
        title_font = Font(size=16, bold=True, color="FFFFFF")
        hdr_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        na_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        nt_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        ws_sum = wb.active
        ws_sum.title = "Summary"

        ws_sum.merge_cells("B2:E3")
        title_cell = ws_sum.cell(row=2, column=2, value="ONT Automation Test Report")
        title_cell.font = title_font
        title_cell.alignment = center_align
        title_cell.fill = hdr_fill

        total = len(cases)
        c_pass = sum(1 for c in cases if c["status"] == "PASS")
        c_fail = sum(1 for c in cases if c["status"] == "FAIL")
        c_na = sum(1 for c in cases if c["status"] == "N/A")
        c_nt = sum(1 for c in cases if c["status"] == "N/T")
        progress = ((total - c_nt) / total) * 100 if total else 0

        info = [
            ("Target Serial", serial),
            ("Target PON / ID", f"{config.get('pon', 'N/A')} / {config.get('ont_id', 'N/A')}"),
            ("OLT Vendor", config.get("olt_vendor", "Unknown")),
            ("OLT Model", config.get("olt_model", "Unknown")),
            ("ONT SW Version", config.get("sw_version", "Unknown")),
            ("VLAN Mode", config.get("vlan_mode", "untagged").capitalize()),
            ("Test Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Total Tests", total),
            ("PASS", c_pass),
            ("FAIL", c_fail),
            ("N/A (Skipped)", c_na),
            ("N/T (Not Tested)", c_nt),
            ("Progress", f"{progress:.1f}%"),
        ]

        row_idx = 5
        for key, val in info:
            c_key = ws_sum.cell(row=row_idx, column=2, value=key)
            c_val = ws_sum.cell(row=row_idx, column=3, value=val)
            c_key.font = bold_font
            c_key.border = thin_border
            c_val.border = thin_border
            if key == "PASS":
                c_val.fill = pass_fill
            elif key == "FAIL":
                c_val.fill = fail_fill
            elif key == "N/A (Skipped)":
                c_val.fill = na_fill
            elif key == "N/T (Not Tested)":
                c_val.fill = nt_fill
            row_idx += 1

        ws_sum.column_dimensions["B"].width = 25
        ws_sum.column_dimensions["C"].width = 30

        ws_det = wb.create_sheet(title="Details")
        headers = ["Test Order", "Test Item", "Command / Script", "Result"]
        for col_num, header in enumerate(headers, 1):
            cell = ws_det.cell(row=1, column=col_num, value=header)
            cell.font = bold_font
            cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            cell.alignment = center_align
            cell.border = thin_border

        for idx, c in enumerate(cases, 2):
            ws_det.cell(row=idx, column=1, value=c["order"]).alignment = center_align
            ws_det.cell(row=idx, column=2, value=c["type"])
            ws_det.cell(row=idx, column=3, value=c["command"])

            res_cell = ws_det.cell(row=idx, column=4, value=c["status"])
            res_cell.alignment = center_align
            if c["status"] == "PASS":
                res_cell.fill = pass_fill
            elif c["status"] == "FAIL":
                res_cell.fill = fail_fill
            elif c["status"] == "N/A":
                res_cell.fill = na_fill
            else:
                res_cell.fill = nt_fill

            for col in range(1, 5):
                ws_det.cell(row=idx, column=col).border = thin_border

        ws_det.column_dimensions["B"].width = 25
        ws_det.column_dimensions["C"].width = 80
        ws_det.column_dimensions["D"].width = 15

        ws_log = wb.create_sheet(title="Logs & Graph")
        ws_log.cell(row=1, column=1, value="Test Item").font = bold_font
        ws_log.cell(row=1, column=2, value="Execution Log").font = bold_font
        ws_log.column_dimensions["A"].width = 20
        ws_log.column_dimensions["B"].width = 100

        log_row = 2
        up_data = {}
        dl_data = {}
        current_mode = "UPLOAD"

        for c in cases:
            ws_log.cell(row=log_row, column=1, value=c["type"]).font = bold_font
            log_lines = (c["log"] or "No log captured.").split("\n")

            for line in log_lines:
                clean_line = line.strip()
                if not clean_line:
                    continue

                clean_line = re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", "", clean_line)
                if clean_line.startswith("="):
                    clean_line = "'" + clean_line

                ws_log.cell(row=log_row, column=2, value=clean_line)

                if c["type"] == "Speed_Test":
                    if "--- STARTING DOWNLOAD TEST ---" in clean_line:
                        current_mode = "DOWNLOAD"
                    elif "--- STARTING UPLOAD TEST ---" in clean_line:
                        current_mode = "UPLOAD"

                    if "[SUM]" in clean_line and "sec" in clean_line and "bits/sec" in clean_line:
                        if "sender" not in clean_line and "receiver" not in clean_line:
                            match = re.search(
                                r"\[SUM\]\s+\d+\.\d+-\s*(\d+\.\d+)\s+sec.*?\s+(\d+(?:\.\d+)?)\s+([KMG])bits/sec",
                                clean_line,
                            )
                            if match:
                                end_time_float = float(match.group(1))
                                t_sec = int(round(end_time_float))
                                val = float(match.group(2))
                                unit = match.group(3)
                                mbps = val * 1000 if unit == "G" else (val if unit == "M" else val / 1000)

                                if current_mode == "UPLOAD":
                                    up_data[t_sec] = max(up_data.get(t_sec, 0), mbps)
                                else:
                                    dl_data[t_sec] = max(dl_data.get(t_sec, 0), mbps)
                log_row += 1
            log_row += 1

        if up_data or dl_data:
            ws_log.cell(row=1, column=4, value="Time (sec)").font = bold_font
            ws_log.cell(row=1, column=5, value="Upload (Mbps)").font = bold_font
            ws_log.cell(row=1, column=6, value="Download (Mbps)").font = bold_font

            all_times = sorted(list(set(list(up_data.keys()) + list(dl_data.keys()))))

            r_idx = 2
            if all_times:
                max_time = max(all_times)
                last_up = None
                last_dl = None

                for t in range(1, max_time + 1):
                    val_up = up_data.get(t, last_up)
                    val_dl = dl_data.get(t, last_dl)

                    ws_log.cell(row=r_idx, column=4, value=t)
                    ws_log.cell(row=r_idx, column=5, value=val_up)
                    ws_log.cell(row=r_idx, column=6, value=val_dl)

                    if val_up is not None:
                        last_up = val_up
                    if val_dl is not None:
                        last_dl = val_dl
                    r_idx += 1

            chart = LineChart()
            chart.title = "iPerf3 Throughput over Time (Up/Down)"
            chart.y_axis.title = "Throughput (Mbps)"
            chart.x_axis.title = "Time (sec)"
            chart.width = 20
            chart.height = 10

            chart.x_axis.tickLblPos = "low"
            chart.y_axis.tickLblPos = "low"
            chart.dispBlanksAs = "span"

            data = Reference(ws_log, min_col=5, max_col=6, min_row=1, max_row=r_idx - 1)
            cats = Reference(ws_log, min_col=4, min_row=2, max_row=r_idx - 1)

            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)

            try:
                for s in chart.series:
                    s.marker = Marker(symbol="circle", size=5)
            except Exception as e:
                logging.warning(f"Could not format chart markers: {e}")

            ws_log.add_chart(chart, "H2")

        wb.save(filename)
        print("\n" + "=" * 60)
        logging.info(f" [✔] Professional Excel Report Generated: {filename} ")
        print("=" * 60 + "\n")

    except Exception as e:
        logging.error(f"Failed to generate Excel report: {e}")
